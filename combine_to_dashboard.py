import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill
import os

VBA_MACRO_CODE = """
Sub ExtractBankLedgers()
    Dim wsBank As Worksheet, wsBankMap As Worksheet
    Dim wsSalesMap As Worksheet, wsPurchMap As Worksheet
    Dim lastRowBank As Long, lastRowBankMap As Long
    Dim i As Long
    Dim partyName As String, natureExp As String, drCr As String
    Dim dictLedgers As Object
    Set dictLedgers = CreateObject("Scripting.Dictionary")
    
    Set wsBank = ThisWorkbook.Sheets("Bank")
    Set wsBankMap = ThisWorkbook.Sheets("Bank_Map")
    Set wsSalesMap = ThisWorkbook.Sheets("Sales_Map")
    Set wsPurchMap = ThisWorkbook.Sheets("Purchase_Map")
    
    Dim checkDict As Object
    Set checkDict = CreateObject("Scripting.Dictionary")
    
    Dim r As Long
    For r = 2 To wsSalesMap.Cells(wsSalesMap.Rows.Count, "A").End(xlUp).Row
        checkDict(wsSalesMap.Cells(r, 1).Value) = True
    Next r
    For r = 2 To wsPurchMap.Cells(wsPurchMap.Rows.Count, "A").End(xlUp).Row
        checkDict(wsPurchMap.Cells(r, 1).Value) = True
    Next r
    
    Dim colParty As Integer, colNature As Integer, colDrCr As Integer
    Dim colWith As Integer, colDep As Integer
    For i = 1 To wsBank.Cells(1, wsBank.Columns.Count).End(xlToLeft).Column
        If wsBank.Cells(1, i).Value = "Party Name" Then colParty = i
        If wsBank.Cells(1, i).Value = "Nature of Expense/Income" Then colNature = i
        If wsBank.Cells(1, i).Value = "Cr/Dr" Then colDrCr = i
        If wsBank.Cells(1, i).Value = "Withdrawals" Then colWith = i
        If wsBank.Cells(1, i).Value = "Deposits" Then colDep = i
    Next i
    
    lastRowBank = wsBank.Cells(wsBank.Rows.Count, colParty).End(xlUp).Row
    
    For i = 2 To lastRowBank
        partyName = Trim(wsBank.Cells(i, colParty).Value)
        natureExp = Trim(wsBank.Cells(i, colNature).Value)
        
        drCr = ""
        If colDrCr > 0 Then
            drCr = UCase(Trim(wsBank.Cells(i, colDrCr).Value))
        ElseIf colWith > 0 And colDep > 0 Then
            If Val(wsBank.Cells(i, colWith).Value) > 0 Then
                drCr = "DR"
            ElseIf Val(wsBank.Cells(i, colDep).Value) > 0 Then
                drCr = "CR"
            End If
        End If
        
        Dim partyAddress As String
        partyAddress = wsBank.Cells(i, colParty).Address(True, True)
        
        Dim natureAddress As String
        natureAddress = wsBank.Cells(i, colNature).Address(True, True)
        
        If partyName <> "" And LCase(partyName) <> "bank account" And LCase(partyName) <> "cash" Then
            If Not checkDict.exists(partyName) Then
                If Not dictLedgers.exists(partyName) Then
                    dictLedgers.Add partyName, IIf(drCr = "DR", "Current Liabilities", "Sundry Debtors") & "|" & partyAddress
                Else
                    If Split(dictLedgers(partyName), "|")(0) <> IIf(drCr = "DR", "Current Liabilities", "Sundry Debtors") Then
                        dictLedgers(partyName) = "CONFLICT|" & Split(dictLedgers(partyName), "|")(1)
                    End If
                End If
            End If
        End If
        
        If natureExp <> "" And LCase(natureExp) <> "cash" Then
            If Not checkDict.exists(natureExp) Then
                If Not dictLedgers.exists(natureExp) Then
                    dictLedgers.Add natureExp, IIf(drCr = "DR", "Indirect Expenses", "Indirect Incomes") & "|" & natureAddress
                Else
                    If Split(dictLedgers(natureExp), "|")(0) <> IIf(drCr = "DR", "Indirect Expenses", "Indirect Incomes") Then
                        dictLedgers(natureExp) = "CONFLICT|" & Split(dictLedgers(natureExp), "|")(1)
                    End If
                End If
            End If
        End If
    Next i
    
    wsBankMap.Cells.Clear
    wsBankMap.Cells(1, 1).Value = "Ledger Name"
    wsBankMap.Cells(1, 2).Value = "Under Group"
    
    Dim key As Variant
    lastRowBankMap = 2
    For Each key In dictLedgers.keys
        wsBankMap.Cells(lastRowBankMap, 1).Formula = "='Bank'!" & Split(dictLedgers(key), "|")(1)
        If Split(dictLedgers(key), "|")(0) = "CONFLICT" Then
            wsBankMap.Cells(lastRowBankMap, 2).Value = ""
            wsBankMap.Cells(lastRowBankMap, 2).Interior.Color = vbYellow
        Else
            wsBankMap.Cells(lastRowBankMap, 2).Value = Split(dictLedgers(key), "|")(0)
        End If
        lastRowBankMap = lastRowBankMap + 1
    Next key
    
    MsgBox "Bank Ledgers Extracted successfully!", vbInformation
End Sub

Sub MergeMasterMap()
    Dim wsMaster As Worksheet
    Dim sheetNames As Variant
    Dim s As Integer
    Dim lastRow As Long, targetRow As Long
    Dim r As Long
    Dim dictUnique As Object
    Set dictUnique = CreateObject("Scripting.Dictionary")
    
    On Error Resume Next
    Set wsMaster = ThisWorkbook.Sheets("MASTER_LEDGER_MAP")
    On Error GoTo 0
    If wsMaster Is Nothing Then
        Set wsMaster = ThisWorkbook.Sheets.Add(After:=ThisWorkbook.Sheets(ThisWorkbook.Sheets.Count))
        wsMaster.Name = "MASTER_LEDGER_MAP"
    Else
        wsMaster.Cells.Clear
    End If
    
    wsMaster.Cells(1, 1).Value = "Ledger Name"
    wsMaster.Cells(1, 2).Value = "Under Group"
    targetRow = 2
    
    sheetNames = Array("Sales_Map", "Purchase_Map", "Bank_Map")
    
    For s = LBound(sheetNames) To UBound(sheetNames)
        Dim ws As Worksheet
        Set ws = ThisWorkbook.Sheets(sheetNames(s))
        lastRow = ws.Cells(ws.Rows.Count, "A").End(xlUp).Row
        
        If lastRow >= 2 Then
            For r = 2 To lastRow
                Dim ledgerName As String
                ledgerName = Trim(ws.Cells(r, 1).Value)
                
                If ledgerName <> "" Then
                    If Not dictUnique.exists(ledgerName) Then
                        dictUnique.Add ledgerName, True
                        ws.Range("A" & r & ":B" & r).Copy
                        wsMaster.Cells(targetRow, 1).PasteSpecial Paste:=xlPasteAll
                        targetRow = targetRow + 1
                    End If
                End If
            Next r
        End If
    Next s
    
    Application.CutCopyMode = False
    wsMaster.Columns("A:B").AutoFit
    MsgBox "Master Ledger Map created! Ready for Tally.", vbInformation
End Sub
"""

def apply_bank_formatting(output_path):
    wb = openpyxl.load_workbook(output_path)
    if 'Bank' not in wb.sheetnames:
        return
    ws = wb['Bank']
    
    headers = [str(cell.value) for cell in ws[1]]
    if 'Confidence' not in headers:
        return
    
    conf_idx = headers.index('Confidence') + 1
    
    green_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    red_fill = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
    
    cyan_fill = PatternFill(start_color="E0FFFF", end_color="E0FFFF", fill_type="solid")
    blue_fill = PatternFill(start_color="C9DAF8", end_color="C9DAF8", fill_type="solid")
    purple_fill = PatternFill(start_color="E4D7F5", end_color="E4D7F5", fill_type="solid")
    orange_fill = PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid")
    
    for row in range(2, ws.max_row + 1):
        conf = ws.cell(row=row, column=conf_idx).value
        fill = None
        if conf == 'High' or conf == 'High-Green':
            fill = green_fill
        elif conf == 'High-Cyan':
            fill = cyan_fill
        elif conf == 'High-Blue':
            fill = blue_fill
        elif conf == 'High-Purple':
            fill = purple_fill
        elif conf == 'High-Orange':
            fill = orange_fill
        elif conf == 'Medium':
            fill = yellow_fill
        elif conf == 'Low':
            fill = red_fill
            
        if fill:
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = fill
                
    wb.save(output_path)

def inject_vba_macro(xlsx_path):
    import win32com.client
    try:
        xl = win32com.client.Dispatch("Excel.Application")
        xl.Visible = False
        xl.DisplayAlerts = False
        
        abs_path = os.path.abspath(xlsx_path)
        wb = xl.Workbooks.Open(abs_path)
        
        # Inject macro
        xlmodule = wb.VBProject.VBComponents.Add(1)
        xlmodule.CodeModule.AddFromString(VBA_MACRO_CODE)
        
        xlsm_path = abs_path.replace(".xlsx", ".xlsm")
        wb.SaveAs(xlsm_path, FileFormat=52) # 52 = xlOpenXMLWorkbookMacroEnabled
        wb.Close()
        xl.Quit()
        print(f"Successfully created Macro-Enabled Dashboard at: {xlsm_path}")
        
        # Cleanup xlsx
        if os.path.exists(abs_path):
            os.remove(abs_path)
            
    except Exception as e:
        print(f"Failed to inject VBA macro. Ensure 'Trust access to the VBA project object model' is checked in Excel Trust Center. Error: {e}")
        try:
            xl.Quit()
        except:
            pass

def create_dashboard(sales_file, purch_file, bank_file, out_path):
    # Read the individual processed outputs
    sales = pd.read_excel(sales_file, sheet_name=None)
    purch = pd.read_excel(purch_file, sheet_name=None)
    bank = pd.read_excel(bank_file)
    
    with pd.ExcelWriter(out_path) as writer:
        sales['RAW_DATA_MASTER'].to_excel(writer, sheet_name='Sales', index=False)
        purch['RAW_DATA_MASTER'].to_excel(writer, sheet_name='Purchases', index=False)
        bank.to_excel(writer, sheet_name='Bank', index=False)
        
        sales['LEDGER_GROUP_MAP'].to_excel(writer, sheet_name='Sales_Map', index=False)
        purch['LEDGER_GROUP_MAP'].to_excel(writer, sheet_name='Purchase_Map', index=False)
        
        pd.DataFrame(columns=['Ledger Name', 'Under Group']).to_excel(writer, sheet_name='Bank_Map', index=False)
        
    # Re-apply OpenPyXL conditional formatting on Bank sheet
    apply_bank_formatting(out_path)
    
    # Inject Macro and save as .xlsm
    inject_vba_macro(out_path)

if __name__ == "__main__":
    sales_file = r"C:\Users\Admin\.gemini\antigravity\scratch\tally_automation\Output\processed_sales.xlsx"
    purch_file = r"C:\Users\Admin\.gemini\antigravity\scratch\tally_automation\Output\processed_purchases.xlsx"
    bank_file = r"C:\Users\Admin\.gemini\antigravity\scratch\tally_automation\Output\processed_bank_v2.xlsx"
    out_path = r"C:\Users\Admin\.gemini\antigravity\scratch\tally_automation\Output\Final_Tally_Dashboard.xlsx"
    create_dashboard(sales_file, purch_file, bank_file, out_path)

