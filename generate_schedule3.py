import openpyxl
import pandas as pd
import math
import os

TB_PATH = r"Infitiminds finalisation\Infinitiminds Trial Balance.xlsx"
FMT_PATH = r"Infitiminds finalisation\Infitiminds Balance Sheet FY 2024-25.xlsx"

GROUP_MAPPING = {
    "Capital Account": "Note 2 : Share Capital ",
    "Loans (Liability)": "Note 4 : Long-term borrowings",
    "Current Liabilities": "Note 7 : Other Current Liabilties",
    "Duties & Taxes": "Note 8 : Short-term Provisions", 
    "Sundry Creditors": "Note 6 : Trade Payable",
    "Other Current Liabilities": "Note 7 : Other Current Liabilties",
    "Fixed Assets": "FA", 
    "Current Assets": "Note 12 : Other Current Assets",
    "Sundry Debtors": "Note 10 : Trade Receivable",
    "Cash-in-Hand": "Note 11 : Cash and cash equivalents",
    "Bank Accounts": "Note 11 : Cash and cash equivalents",
    "Sales Accounts": "Note 13 : REVENUE FROM OPERATIONS",
    "Purchase Accounts": "Note 14 : COST OF RENDERING SERVICES",
    "Direct Expenses": "Note 14 : COST OF RENDERING SERVICES",
    "Indirect Expenses": "Note 16 : Other Expenses",
}

def clean_val(v):
    if v is None: return 0.0
    try:
        return float(v)
    except:
        return 0.0

def parse_tb():
    wb = openpyxl.load_workbook(TB_PATH, data_only=True)
    ws = wb.active
    
    rows = []
    start_row = 1
    for r in range(1, 20):
        val = str(ws.cell(row=r, column=1).value).strip()
        if "Particulars" in val or "Capital Account" in val:
            start_row = r
            if "Particulars" in val: start_row += 2
            break
            
    for r in range(start_row, ws.max_row + 1):
        cell = ws.cell(row=r, column=1)
        name_raw = str(cell.value)
        if not cell.value or name_raw.strip() == "" or name_raw == "None" or "Grand Total" in name_raw:
            continue
            
        indent = cell.alignment.indent or 0.0
        spaces = float(indent)
        name = name_raw.strip()
        
        op = clean_val(ws.cell(row=r, column=2).value)
        dr = clean_val(ws.cell(row=r, column=3).value)
        cr = clean_val(ws.cell(row=r, column=4).value)
        cl = clean_val(ws.cell(row=r, column=5).value)
        
        if cl == 0:
            continue
            
        bal_type = "UNKNOWN"
        if math.isclose(op + dr - cr, cl, abs_tol=1.0) or math.isclose(-op + dr - cr, cl, abs_tol=1.0) or math.isclose(dr - cr, cl, abs_tol=1.0):
            bal_type = "DR"
        elif math.isclose(op - dr + cr, cl, abs_tol=1.0) or math.isclose(-op - dr + cr, cl, abs_tol=1.0) or math.isclose(cr - dr, cl, abs_tol=1.0):
            bal_type = "CR"
            
        rows.append({
            'row': r,
            'spaces': spaces,
            'name': name,
            'op': op, 'dr': dr, 'cr': cr, 'cl': cl,
            'type': bal_type
        })
        
    ledgers = []
    current_parents = {}
    
    for i, row in enumerate(rows):
        is_group = False
        if i < len(rows) - 1:
            if rows[i+1]['spaces'] > row['spaces']:
                is_group = True
                
        current_parents[row['spaces']] = row['name']
        
        if not is_group:
            parent_group = "Unknown"
            for s in sorted(current_parents.keys(), reverse=True):
                if s < row['spaces']:
                    parent_group = current_parents[s]
                    break
                    
            if parent_group == "Unknown":
                parent_group = row['name']
                
            mapped_note = GROUP_MAPPING.get(parent_group, "UNMAPPED")
            
            if "Salary" in row['name'] or "Employee" in row['name']:
                mapped_note = "Note 15 : Employee Benefits Expense"
            if "Tds Receivable" in row['name']:
                mapped_note = "Note 12 : Other Current Assets"
            
            ledgers.append({
                'Parent Group': parent_group,
                'Ledger': row['name'],
                'Closing Balance': row['cl'],
                'Dr/Cr': row['type'],
                'Mapped Note': mapped_note
            })
            
    return pd.DataFrame(ledgers)

if __name__ == "__main__":
    df = parse_tb()
    print("Parsed TB:")
    print(df.to_string())
    
    unmapped = df[df['Mapped Note'] == 'UNMAPPED']
    if not unmapped.empty:
        print("\nWARNING - UNMAPPED LEDGERS:")
        print(unmapped.to_string())
